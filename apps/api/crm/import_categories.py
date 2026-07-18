"""Import / export des catégories de documents via un fichier Excel.

Alimente la table `categories` (cf. crm/db.py v8 et v16) à partir d'un
classeur .xlsx contenant (Nom, Couleur, Icone, Ordre, Message WhatsApp).

Idempotent : relancer le même fichier met à jour les catégories existantes
(rapprochées par leur Nom) ou crée les nouvelles.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Optional

from . import repo
from .backup import backup_db
from .db import connect, default_db_path
from .repo import slugify

try:
    from src import ui
except Exception:  # noqa: BLE001
    ui = None  # type: ignore[assignment]


# Mots-clés d'en-tête reconnus
_NOM_KEYS = {
    "nom",
    "noms",
    "name",
    "names",
    "libelle",
    "libelles",
    "label",
    "labels",
    "designation",
}
_COULEUR_KEYS = {"couleur", "couleurs", "color", "colors", "hex"}
_ICONE_KEYS = {"icone", "icones", "icon", "icons"}
_ORDRE_KEYS = {"ordre", "ordres", "sort_order", "sort", "tri"}
_WHATSAPP_KEYS = {
    "message_whatsapp",
    "whatsapp",
    "whatsapp_message",
    "message_wa",
    "wa_message",
    "message",
}


@dataclass
class ImportSummary:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = None  # type: ignore[assignment]

    def __post_init__(self) -> None:
        if self.errors is None:
            self.errors = []


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401

        return openpyxl
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "Le module « openpyxl » (fichiers .xlsx) est introuvable. "
            "En developpement : pip install openpyxl."
        ) from exc


def _clean_string(value: object) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def _parse_ordre(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, bool):
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    s = str(value).strip()
    if not s or not re.fullmatch(r"\-?\d+", s):
        return 0
    return int(s)


def _detect_columns(header: tuple) -> Optional[dict[str, Optional[int]]]:
    cols: dict[str, Optional[int]] = {
        "nom": None,
        "couleur": None,
        "icone": None,
        "ordre": None,
        "whatsapp": None,
    }
    for i, cell in enumerate(header):
        slug = slugify(str(cell)) if cell is not None else ""
        if not slug:
            continue
        if cols["nom"] is None and slug in _NOM_KEYS:
            cols["nom"] = i
        elif cols["couleur"] is None and slug in _COULEUR_KEYS:
            cols["couleur"] = i
        elif cols["icone"] is None and slug in _ICONE_KEYS:
            cols["icone"] = i
        elif cols["ordre"] is None and slug in _ORDRE_KEYS:
            cols["ordre"] = i
        elif cols["whatsapp"] is None and slug in _WHATSAPP_KEYS:
            cols["whatsapp"] = i

    if cols["nom"] is None:
        return None  # aucune entête reconnue
    return cols


# Une ligne lue : (no_ligne, nom, couleur, icone, ordre, whatsapp_message)
_Row = tuple[int, str, Optional[str], Optional[str], int, Optional[str]]


def _read_rows(path: Path, feuille: Optional[str]) -> list[_Row]:
    openpyxl = _require_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if feuille:
            if feuille not in wb.sheetnames:
                raise RuntimeError(
                    f"Feuille « {feuille} » introuvable. "
                    f"Feuilles disponibles : {', '.join(wb.sheetnames)}"
                )
            ws = wb[feuille]
        else:
            ws = wb.active

        all_rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not all_rows:
        return []

    cols = _detect_columns(all_rows[0])
    if cols is not None:
        data_rows = all_rows[1:]
        start_line = 2
        i_nom = cols["nom"]
        i_coul = cols["couleur"]
        i_ico = cols["icone"]
        i_ord = cols["ordre"]
        i_wa = cols["whatsapp"]
    else:
        # Mode positionnel par défaut : 1=Nom, 2=Couleur, 3=Icone, 4=Ordre, 5=WhatsApp
        data_rows = all_rows
        start_line = 1
        i_nom, i_coul, i_ico, i_ord, i_wa = 0, 1, 2, 3, 4

    out: list[_Row] = []
    for offset, row in enumerate(data_rows):
        line_no = start_line + offset

        def _at(idx: Optional[int]):
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        nom = _clean_string(_at(i_nom))
        if nom is None:
            continue  # Ligne ignorée si le nom est manquant

        out.append(
            (
                line_no,
                nom,
                _clean_string(_at(i_coul)),
                _clean_string(_at(i_ico)),
                _parse_ordre(_at(i_ord)),
                _clean_string(_at(i_wa)),
            )
        )
    return out


def _import_rows(conn, rows: list[_Row], dry_run: bool) -> ImportSummary:
    summary = ImportSummary()
    for line_no, nom, couleur, icone, ordre, whatsapp_message in rows:
        try:
            existing = repo.get_category(conn, nom)
            if existing is not None:
                existing.couleur = couleur or existing.couleur
                existing.icone = icone or existing.icone
                existing.sort_order = ordre
                existing.whatsapp_message = whatsapp_message
                if not dry_run:
                    repo.upsert_category(conn, existing)
                summary.updated += 1
            else:
                if not dry_run:
                    repo.upsert_category(
                        conn,
                        repo.Category(
                            nom=nom,
                            couleur=couleur,
                            icone=icone,
                            sort_order=ordre,
                            whatsapp_message=whatsapp_message,
                        ),
                    )
                summary.created += 1
        except Exception as exc:  # noqa: BLE001
            summary.errors.append(f"Ligne {line_no} ({nom}) : {exc}")
            summary.skipped += 1
    return summary


def import_categories(
    path: Path, feuille: Optional[str] = None, dry_run: bool = False, conn=None
) -> ImportSummary:
    rows = _read_rows(path, feuille)
    own = conn is None
    if own:
        conn = connect()
    try:
        return _import_rows(conn, rows, dry_run)
    finally:
        if own:
            conn.close()


_EXPORT_HEADERS = ["Nom", "Couleur", "Icone", "Ordre", "Message WhatsApp"]


def _build_workbook(categories: list[repo.Category]):
    openpyxl = _require_openpyxl()
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Categories"
    ws.append(_EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for c in categories:
        ws.append(
            [
                c.nom,
                c.couleur or "",
                c.icone or "",
                c.sort_order,
                c.whatsapp_message or "",
            ]
        )
    for col, width in zip("ABCDE", (24, 12, 12, 10, 48)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    return wb


def export_categories(path: Path, conn=None) -> int:
    own = conn is None
    if own:
        conn = connect()
    try:
        categories = repo.list_categories(conn)
        _build_workbook(categories).save(path)
        return len(categories)
    finally:
        if own:
            conn.close()


def export_categories_bytes(conn=None) -> tuple[bytes, int]:
    own = conn is None
    if own:
        conn = connect()
    try:
        categories = repo.list_categories(conn)
        buf = BytesIO()
        _build_workbook(categories).save(buf)
        return buf.getvalue(), len(categories)
    finally:
        if own:
            conn.close()


def write_template(path: Path) -> None:
    openpyxl = _require_openpyxl()
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Categories"
    ws.append(_EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDE", (24, 12, 12, 10, 48)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="crm.import_categories",
        description="Importe / exporte la configuration des catégories (Nom, Couleur, Icone, Ordre, Message WhatsApp) en .xlsx.",
    )
    parser.add_argument("fichier", nargs="?", help="Classeur .xlsx à importer.")
    parser.add_argument("--feuille", help="Nom de la feuille (défaut : la active).")
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Simulation : affiche ce qui serait fait sans rien écrire en base.",
    )
    parser.add_argument(
        "--export",
        metavar="FICHIER.xlsx",
        help="Exporte les catégories vers ce fichier, puis quitte.",
    )
    parser.add_argument(
        "--modele",
        metavar="FICHIER.xlsx",
        help="Génère un fichier modèle vierge, puis quitte.",
    )
    args = parser.parse_args(argv)

    def _say(level: str, text: str) -> None:
        fn = getattr(ui, level, None) if ui is not None else None
        if fn is not None:
            fn(text)
        else:
            print(text)

    if args.export:
        try:
            n = export_categories(Path(args.export))
        except Exception as exc:  # noqa: BLE001
            _say("error", f"Echec de l'export : {exc}")
            return 1
        _say("success", f"Export terminé : {n} catégorie(s) -> {args.export}")
        return 0

    if args.modele:
        try:
            write_template(Path(args.modele))
        except Exception as exc:  # noqa: BLE001
            _say("error", f"Echec de génération du modèle : {exc}")
            return 1
        _say("success", f"Modèle créé : {args.modele}")
        return 0

    if not args.fichier:
        parser.error("indiquez un fichier .xlsx à importer (ou --export / --modele).")

    path = Path(args.fichier)
    if not path.is_file():
        _say("error", f"Fichier introuvable : {path}")
        return 1

    if ui is not None:
        ui.banner("Import des catégories de documents", str(path))

    if not args.dry_run:
        dest = backup_db()
        if dest is not None:
            _say("note", f"Sauvegarde préalable : {dest}")

    try:
        summary = import_categories(path, feuille=args.feuille, dry_run=args.dry_run)
    except RuntimeError as exc:
        _say("error", str(exc))
        return 1

    if args.dry_run:
        _say("warn", "SIMULATION : aucune modification écrite en base.")

    if summary.errors:
        if ui is not None:
            ui.section("Erreurs / Lignes ignorées")
        for msg in summary.errors:
            _say("warn", msg)

    if ui is not None:
        print()
        ui.success("Import terminé." if not args.dry_run else "Simulation terminée.")
        ui.stat("Catégories créées", summary.created, accent=ui.GREEN)
        ui.stat("Catégories mises à jour", summary.updated, accent=ui.GREEN)
    else:
        print("Import terminé." if not args.dry_run else "Simulation terminée.")
        print(f"  Catégories créées : {summary.created}")
        print(f"  Catégories mises à jour : {summary.updated}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
